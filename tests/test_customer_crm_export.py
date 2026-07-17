import unittest
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from dashboard.crm_export import build_customer_crm_workbook


class CustomerCrmExportTests(unittest.TestCase):
    def test_workbook_contains_formatted_customer_rows(self):
        customer = SimpleNamespace(
            id=7,
            name="테스트 고객",
            phone="01012345678",
            email="test@example.com",
            booking_restricted=False,
        )
        rows = [{
            "customer": customer,
            "segment_label": "잠재 VIP",
            "segment_reason": "3 completed visits",
            "completed_count": 3,
            "booking_count": 5,
            "total_revenue": 210000,
            "average_ticket": 70000,
            "last_visit_at": datetime(2026, 7, 10, 14, 0),
            "next_booking_at": datetime(2026, 8, 1, 15, 30),
            "cancelled_count": 1,
            "cancellation_rate": 20.0,
            "no_show_count": 0,
            "no_show_rate": 0.0,
            "preferred_service": "젤 네일",
            "preferred_staff": "Julie",
        }]

        stream = build_customer_crm_workbook(
            rows,
            query="테스트",
            segment="potential_vip",
            status="normal",
            sort_key="revenue",
        )
        workbook = load_workbook(stream)
        sheet = workbook["고객 현황"]

        self.assertEqual(sheet["A1"].value, "Salon De Nature 고객 CRM 현황")
        self.assertEqual(sheet["B6"].value, "테스트 고객")
        self.assertEqual(sheet["I6"].value, 210000)
        self.assertEqual(sheet["I6"].number_format, '"₩"#,##0')
        self.assertEqual(sheet["N6"].value, 0.2)
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertIn("CustomerCRMTable", sheet.tables)


if __name__ == "__main__":
    unittest.main()
